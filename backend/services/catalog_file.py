"""Leer un catálogo desde un archivo (Excel .xlsx o .csv) que nos manda un vendedor.

Lo usa el admin para subir de golpe el catálogo de un vendedor que no tiene web propia
(por ejemplo, el archivo que usó para Importamatic de Todocolección o una lista hecha a mano).

Las columnas se reconocen por su nombre, en español o inglés y sin importar mayúsculas ni tildes:
  título / nombre / artículo       (obligatoria)
  precio                           (obligatoria para publicar; se puede completar en la revisión)
  descripción
  categoría / sección
  estado                           (nuevo, usado, restaurado)
  referencia / ref / código / sku  (sirve para emparejar las fotos por nombre de archivo)
  foto / fotos / imagen / imagen1…  (nombres de archivo o direcciones https, separadas por comas)
"""

import csv
import io
import re
import unicodedata
from typing import Optional

from services.catalog_import import DESCRIPTION_MAX, MAX_PHOTOS, TITLE_MAX, parse_price, strip_html

MAX_ROWS = 1000
MAX_FILE_BYTES = 10 * 1024 * 1024


class CatalogFileError(ValueError):
    """Error con un mensaje apto para mostrar en el admin."""


def _norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


FIELDS = {
    "title": ("titulo", "title", "nombre", "articulo", "producto", "nombre del articulo", "titulo del lote", "lote"),
    "price": ("precio", "price", "pvp", "importe", "precio venta", "precio de venta", "precio eur"),
    "description": ("descripcion", "description", "desc", "detalle", "detalles", "texto", "descripcion del lote"),
    "category": ("categoria", "category", "seccion", "tipo", "familia"),
    "condition": ("estado", "condition", "conservacion", "estado de conservacion"),
    "ref": ("referencia", "ref", "codigo", "cod", "sku", "id", "numero", "n", "num", "referencia interna"),
}
PHOTO_PREFIXES = ("foto", "fotos", "imagen", "imagenes", "image", "images", "img", "photo", "photos", "picture")


def _field_for(header: str) -> Optional[str]:
    h = _norm(header)
    if not h:
        return None
    for field, names in FIELDS.items():
        if h in names:
            return field
    first = h.split(" ")[0]
    if first in PHOTO_PREFIXES or re.match(r"^(foto|imagen|image|img|photo)\s*\d+$", h):
        return "photos"
    # Nombres más largos: "Título del artículo", "Precio (€)"...
    for field, names in FIELDS.items():
        if field != "ref" and first in names:
            return field
    return None


def _condition(value) -> str:
    v = _norm(value)
    if not v:
        return ""
    if "restaur" in v:
        return "restaurado"
    if v.startswith("nuev") or v in {"new", "a estrenar", "sin usar"}:
        return "nuevo"
    return "usado"


def _price(value) -> Optional[float]:
    """Como parse_price, pero "2.000" es dos mil (formato español), no dos."""
    if isinstance(value, str):
        compact = re.sub(r"[\s€]|eur(os)?", "", value.strip(), flags=re.I)
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", compact):
            value = compact.replace(".", "")
    return parse_price(value)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_rows(filename: str, data: bytes) -> list[list]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover
            raise CatalogFileError("No se pueden leer archivos Excel ahora mismo. Guárdalo como CSV.") from exc
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise CatalogFileError("No se ha podido abrir el Excel. Prueba a guardarlo de nuevo como .xlsx o .csv.") from exc
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if name.endswith(".xls"):
        raise CatalogFileError("Ese Excel es del formato antiguo (.xls). Ábrelo y guárdalo como .xlsx o .csv.")

    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
    except csv.Error:
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _find_header(rows: list[list]) -> tuple[int, dict[int, str]]:
    """Busca la fila de cabecera entre las 10 primeras (a veces hay un título encima)."""
    best = (-1, {})
    for i, row in enumerate(rows[:10]):
        mapping = {}
        for j, header in enumerate(row):
            field = _field_for(_cell(header))
            if field and (field == "photos" or field not in mapping.values()):
                mapping[j] = field
        if "title" in mapping.values() and len(mapping) > len(best[1]):
            best = (i, mapping)
    if best[0] < 0:
        raise CatalogFileError(
            "No encuentro la columna del título. La primera fila debe tener los nombres de las columnas "
            "(por ejemplo: Título, Precio, Descripción, Fotos)."
        )
    return best


def _split_photos(value: str) -> list[str]:
    parts = re.split(r"[,;\n|]+|\s+(?=https?://)", value)
    return [p.strip().strip('"') for p in parts if p.strip()]


def parse_catalog_file(filename: str, data: bytes) -> dict:
    if len(data) > MAX_FILE_BYTES:
        raise CatalogFileError("El archivo pesa demasiado (máximo 10 MB). Las fotos van aparte, no dentro del Excel.")
    rows = _read_rows(filename, data)
    header_row, mapping = _find_header(rows)

    items, skipped = [], 0
    for row in rows[header_row + 1:]:
        values: dict[str, str] = {}
        photos: list[str] = []
        for j, field in mapping.items():
            cell = _cell(row[j]) if j < len(row) else ""
            if not cell:
                continue
            if field == "photos":
                photos.extend(_split_photos(cell))
            elif field == "price":
                values["price"] = cell if not isinstance(row[j], (int, float)) else row[j]
            else:
                values[field] = cell
        title = strip_html(values.get("title", ""))[:TITLE_MAX]
        if len(title) < 3:
            if any(values.values()) or photos:
                skipped += 1
            continue
        clean_photos = []
        for p in photos:
            if p.startswith("//"):
                p = "https:" + p
            if p not in clean_photos:
                clean_photos.append(p)
        items.append({
            "title": title,
            "description": strip_html(values.get("description", ""))[:DESCRIPTION_MAX] or None,
            "price": _price(values.get("price")),
            "category": values.get("category") or None,
            "condition": _condition(values.get("condition")),
            "ref": (values.get("ref") or "")[:100] or None,
            "photos": clean_photos[: MAX_PHOTOS * 2],
        })
        if len(items) >= MAX_ROWS:
            break

    if not items:
        raise CatalogFileError("El archivo no tiene ningún artículo con título debajo de la cabecera.")
    return {
        "items": items,
        "skipped": skipped,
        "columns": sorted(set(mapping.values())),
    }
